"""Sync optional commerce columns from the Zylra catalog workbook.

Supported headers (case-insensitive): SKU, Price, Sale Price, Currency, Quantity.
The current V1 workbook has no commerce columns, so the generated overlay will
contain no invented prices or stock counts until those columns are added.
"""
from pathlib import Path
import json
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "catalog-source" / "zylra_product_catalog_v1.xlsx"
OUTPUT = ROOT / "data" / "product-commerce.json"

def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

def number(value):
    if value in (None, ""): return None
    try: return float(str(value).replace(",", "").strip())
    except ValueError: return None

wb = openpyxl.load_workbook(SOURCE, data_only=True)
ws = wb["Product catalog"]
headers = {key(cell.value): i for i, cell in enumerate(ws[1], start=1)}
required = headers.get("sku")
if not required:
    raise SystemExit("SKU column is required")

def get(row, *names):
    for name in names:
        idx = headers.get(key(name))
        if idx: return row[idx - 1].value
    return None

result = {}
for row in ws.iter_rows(min_row=2):
    sku = str(get(row, "SKU") or "").strip()
    if not sku: continue
    result[sku] = {
        "price": number(get(row, "Price")),
        "salePrice": number(get(row, "Sale Price", "SalePrice")),
        "currency": str(get(row, "Currency") or "INR").strip() or "INR",
        "inventoryQuantity": int(number(get(row, "Quantity", "Inventory Quantity", "Stock", "Inventory")) or 0) if any(headers.get(key(n)) for n in ["Quantity", "Inventory Quantity", "Stock", "Inventory"]) else None,
    }
OUTPUT.write_text(json.dumps(result, indent=2), encoding="utf-8")
print(f"Wrote {len(result)} SKU records to {OUTPUT}")
print("Commerce columns detected:", ", ".join(k for k in ["Price", "Sale Price", "Currency", "Quantity"] if headers.get(key(k))) or "none")
